"""Author explicitly synthetic F01-F04 inputs; never writes the original references."""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas


ROOT = Path(__file__).with_name("cases") / "studio-tables" / "inputs"


def workbook(name, headers, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = "Synthetic data"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173D54")
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 28
    path = ROOT / name
    book.save(path)
    observed = list(load_workbook(path, read_only=True).active.values)
    assert observed == [tuple(headers), *map(tuple, rows)], name


def main():
    ROOT.mkdir(parents=True, exist_ok=True)
    workbook("registration-a.xlsx", ["姓名", "邮箱", "部门", "报名日期"], [
        ["张青", "zhang@example.test", "研发", "2026/10/01"],
        ["李明", "li.one@example.test", "销售", "2026/10/02"],
        ["王宁", None, "研发", "2026/10/03"],
    ])
    workbook("registration-b.xlsx", ["报名日期", "部门", "邮箱", "姓名"], [
        ["2026-10-01", "研发", "ZHANG@EXAMPLE.TEST", "张青"],
        ["2026-10-04", "销售", "li.two@example.test", "李明"],
        ["2026-10-03", "研发", None, "王宁"],
    ])
    workbook("dirty-data.xlsx", ["id", "date", "sku", "value"], [
        ["A1", "2026/10/01", "abc", 0],
        ["A1", "2026-10-01", "AbC", 0],
        ["A2", "2026-02-30", "abc", None],
        ["A3", None, "XYZ", 10],
        ["A4", "2026-02-29", "efg", 8],
        ["A5", "2024/02/29", "hij", 2],
    ])
    invoices = [
        ("DEMO-A", "2026-10-11", "123.45", "DEMO-TAX-A"),
        ("DEMO-B", "2026-10-12", "0.00", "NOT PROVIDED"),
        ("DEMO-A", "2026-10-11", "123.45", "DEMO-TAX-A"),
    ]
    path = ROOT / "invoices.pdf"
    pdf = Canvas(str(path), pagesize=A4, pageCompression=0, invariant=1)
    for index, (number, date, amount, tax_id) in enumerate(invoices, 1):
        pdf.setFont("Helvetica-Bold", 16)
        pdf.drawString(52, 780, "SYNTHETIC INVOICE - NOT VALID FOR PAYMENT")
        pdf.setFont("Helvetica", 13)
        for line, text in enumerate([
            f"Page: {index}", f"Invoice ID: {number}", f"Date: {date}",
            f"Total amount (CNY): {amount}", f"Tax ID: {tax_id}",
        ]):
            pdf.drawString(52, 730 - line * 32, text)
        pdf.showPage()
    pdf.save()
    observed = PdfReader(path)
    assert len(observed.pages) == 3
    for page, expected in zip(observed.pages, invoices):
        assert all(value in page.extract_text() for value in expected)
    print("Authored and reopened 3 synthetic workbooks and one three-page PDF.")


if __name__ == "__main__":
    main()
